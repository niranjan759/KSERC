"""The KSERC Truing Up Order house style, for drafting the next order in Word.

Measured from the 2023-24 Truing Up Order PDF (A4, 237 pages):
  page              A4; text from 72 pt to 527 pt across (margins 25.4 mm /
                    about 24 mm); page number at the foot, centred, Arial 11
  body text         Arial 12, justified, about 14 pt between lines
  numbered clause   "2.5" at the margin, its text hanging at 108 pt (36 pt in)
  chapter heading   Arial Bold 14, centred ("CHAPTER-2" / "TRUING UP OF ...")
  section heading   Arial Bold 14 / 13 at 85 pt ("Introduction", "Hydro
                    generation", "Analysis and decision of the Commission")
  minor heading     Arial Bold 12, aligned with the clause text
  conclusion        Arial Bold Italic 12 ("... the Commission hereby approve")
  quoted Regulation Arial Italic 11, indented (to about 143 pt)
  table caption     "Table 2.2" then its title, each centred, Arial 10
  table text        Arial 9, single-line borders, column widths as printed

Each text block of the previous order carries the style it was printed in
(extract_sbu_g._line_style / assemble_blocks), so a block is set the way the
same block was set last year; these constants are the fallback for text
that has no such record (wording written for the new order)."""
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.shared import Mm, Pt

FONT = "Arial"
BODY_PT, QUOTE_PT, CAPTION_PT, TABLE_PT = 12, 11, 10, 9
LEFT_MARGIN_PT = 72.0          # where the text block starts on the page
TEXT_WIDTH_PT = 455.0          # 72 pt .. 527 pt
HANG_PT = 36.0                 # clause text hangs 36 pt in from its number
CLAUSE_START = re.compile(r'^(\d{1,2}\.\d{1,3}(?:\.\d+)?)\s+')
MARKS = re.compile(r'(\[\[.*?\]\]|\{\{.*?\}\}|<<.*?>>)')


def _mark(part):
    """(text shown, highlight colour or None) for a piece of MARKS-split
    text: [[note]] -> yellow "[note]", {{figure}} -> yellow figure, <<figure>>
    -> green figure (moved on to this year's, see order_draft)."""
    if part.startswith("[[") and part.endswith("]]"):
        return "[" + part[2:-2] + "]", "yellow"
    if part.startswith("{{") and part.endswith("}}"):
        return part[2:-2], "yellow"
    if part.startswith("<<") and part.endswith(">>"):
        return part[2:-2], "green"
    return part, None
NUMBER_CELL = re.compile(r'^\(?[-+]?[\d,]*\.?\d+\)?%?$|^-$')
TOTAL_ROW = re.compile(r'^(sub\s*-?\s*total|grand\s+total|total)\b', re.I)


def _set_font(run_or_style, size=None, bold=None, italic=None):
    font = run_or_style.font
    font.name = FONT
    rpr = run_or_style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    if size is not None:
        font.size = Pt(size)
    if bold:
        font.bold = True
    if italic:
        font.italic = True


def new_document():
    """An empty Word document set up like the orders: A4, the orders'
    margins, Arial 12 justified body text, page number centred at the foot."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(25.4), Mm(24.2)
    sec.top_margin, sec.bottom_margin = Mm(25.4), Mm(25.4)
    sec.footer_distance = Mm(12)

    normal = doc.styles["Normal"]
    _set_font(normal, BODY_PT)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    _set_font(run, 11)
    return doc


def write_runs(p, text, size=None, bold=False, italic=False):
    """Text into a paragraph: [[note]] becomes a yellow "[note]" placeholder,
    {{figure}} a yellow-highlighted figure, <<figure>> a green one (_mark)."""
    for part in MARKS.split(text or ""):
        if not part:
            continue
        shown, mark = _mark(part)
        run = p.add_run(shown)
        # the font family comes from the Normal style (Arial, new_document);
        # naming it again on each of 20,000 runs cost most of a minute
        if size is not None:
            run.font.size = Pt(size)
        if bold:
            run.font.bold = True
        if italic:
            run.font.italic = True
        if mark:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW if mark == "yellow" else WD_COLOR_INDEX.BRIGHT_GREEN
    return p


def chapter_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    return write_runs(p, text, 14, bold=True)


def heading(doc, text, style=None):
    """A section or minor heading, at the size and indent it had last year
    (Bold 13 at 85 pt for "Hydro generation"; Bold 12 at 108 pt for a minor
    one); Bold 13 at 85 pt when there's no record."""
    style = style or {}
    size = style.get("size") or 13
    indent = max((style.get("x0") or 85.0) - LEFT_MARGIN_PT, 0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.left_indent = Pt(min(indent, HANG_PT))
    pf.keep_with_next = True
    pf.space_before = Pt(6)
    return write_runs(p, text, size, bold=True, italic=bool(style.get("italic")))


def paragraph(doc, text, style=None):
    """A paragraph set as the same paragraph was last year: a numbered clause
    with its text hanging 36 pt in; an unnumbered one at its old indent; a
    quoted Regulation in Italic 11; the Commission's conclusion in Bold
    Italic."""
    style = style or {}
    size = style.get("size") or BODY_PT
    italic, bold = bool(style.get("italic")), bool(style.get("bold"))
    p = doc.add_paragraph()
    pf = p.paragraph_format
    m = CLAUSE_START.match(text or "")
    if m:
        pf.left_indent = Pt(HANG_PT)
        pf.first_line_indent = Pt(-HANG_PT)
        pf.tab_stops.add_tab_stop(Pt(HANG_PT))
        write_runs(p, m.group(1) + "\t", size)
        text = text[m.end():]
    else:
        first = (style.get("x0") or LEFT_MARGIN_PT + HANG_PT) - LEFT_MARGIN_PT
        later = (style.get("indent") or style.get("x0") or LEFT_MARGIN_PT + HANG_PT) - LEFT_MARGIN_PT
        pf.left_indent = Pt(max(later, 0))
        if abs(first - later) > 2:
            pf.first_line_indent = Pt(first - later)
    return write_runs(p, text, size, bold=bold, italic=italic)


def caption(doc, table_no, title):
    """ "Table 2.2" and its title, each on its own centred line."""
    p = None
    for text in ([f"Table {table_no}"] if table_no else []) + ([title] if title else []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.keep_with_next = True
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        write_runs(p, text, CAPTION_PT)
    if p is not None:
        p.paragraph_format.space_after = Pt(4)


def _kept_columns(header, layout, rows):
    """Columns worth printing: anything with text in a heading or any row.
    The extraction carries hair-thin spacer columns (5 pt wide, never any
    text) that the printed table doesn't show."""
    ncols = max([len(header)] + [len(r) for r in rows])
    used = set()
    for j in range(ncols):
        if j < len(header) and (header[j] or "").strip():
            used.add(j)
        if any(j < len(r) and (r[j] or "").strip() for r in rows):
            used.add(j)
    for hr in (layout or {}).get("header_rows", []):
        for c in hr:
            if (c.get("text") or "").strip():
                span = range(c["col"], c["col"] + c.get("colspan", 1))
                if not used & set(span):
                    used.add(c["col"])
    return sorted(used)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


TABLE_STYLES = {"text": "KSERC Table", "heading": "KSERC Table Heading", "figure": "KSERC Table Figure"}


def _table_styles(doc):
    """Paragraph styles for table cells - Arial 9, no spacing; headings bold
    and centred, figures right-aligned. One style reference per cell instead
    of a font size, spacing and alignment on every cell keeps a 250-table
    order quick to build."""
    from docx.enum.style import WD_STYLE_TYPE
    out = {}
    for kind, name in TABLE_STYLES.items():
        try:
            out[kind] = doc.styles[name]
            continue
        except KeyError:
            pass
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.font.size = Pt(TABLE_PT)
        pf = st.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.0
        pf.alignment = {"text": WD_ALIGN_PARAGRAPH.LEFT, "heading": WD_ALIGN_PARAGRAPH.CENTER,
                        "figure": WD_ALIGN_PARAGRAPH.RIGHT}[kind]
        if kind == "heading":
            st.font.bold = True
        out[kind] = st
    return out


def _cell_text(cell, text, style, size=None, bold=False, italic=False):
    """A cell's text, written straight into its XML: the cell's paragraph
    gets the style's id and one run per piece of text (a highlighted run for
    a [[ ]] / {{ }} mark). python-docx's own setters check element order on
    every call, which for the 20,000 cells of a full order cost most of the
    build time."""
    from lxml import etree
    p_el = cell._tc.p_lst[0]
    ppr = etree.SubElement(p_el, qn("w:pPr"))
    p_el.insert(0, ppr)
    etree.SubElement(ppr, qn("w:pStyle")).set(qn("w:val"), style.style_id)
    for part in MARKS.split(text or ""):
        if not part:
            continue
        shown, mark = _mark(part)
        r = etree.SubElement(p_el, qn("w:r"))
        if bold or italic or size or mark:
            rpr = etree.SubElement(r, qn("w:rPr"))  # children in the schema's order
            if bold:
                etree.SubElement(rpr, qn("w:b"))
            if italic:
                etree.SubElement(rpr, qn("w:i"))
            if size:
                etree.SubElement(rpr, qn("w:sz")).set(qn("w:val"), str(int(round(size * 2))))
            if mark:
                etree.SubElement(rpr, qn("w:highlight")).set(qn("w:val"), mark)
        t = etree.SubElement(r, qn("w:t"))
        t.text = shown
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


CELL_PADDING_PT = 12.0  # Word's default cell margins, 0.08" each side
PROSE_SHARE = 0.45      # of the text width, the most a column of running text asks for


def _balanced_widths(widths, cols, header_rows, rows, size):
    """The printed column widths, corrected where the text wouldn't fit. The
    widths measured off the page can put a column of labels or running text
    in a sliver (a label printed in a column the extraction took as a 5 pt
    spacer, Table 5.90; a stakeholder's comments in Chapter 1), and in a
    fixed-layout table Word then breaks every word letter by letter - one
    table ran to fourteen pages. So no column is narrower than its longest
    word, a column holding sentences asks for room up to PROSE_SHARE of the
    page, and what that needs is taken from the columns with room to spare.
    A table whose widths already fit is left as printed."""
    char = 5.0 * size / TABLE_PT  # an Arial character, a little on the wide side
    pos = {j: k for k, j in enumerate(cols)}
    texts = [[] for _ in cols]
    for row in rows:
        for j in cols:
            if j < len(row) and row[j]:
                texts[pos[j]].append(re.sub(r'\[\[\s*\]\]', '[ ]', re.sub(r'\{\{|\}\}', '', row[j])))
    for hr in header_rows:
        for c in hr:
            kept = [j for j in range(c["col"], c["col"] + c.get("colspan", 1)) if j in pos]
            if len(kept) == 1 and c.get("text"):
                texts[pos[kept[0]]].append(c["text"])
    floor, wish = [], []
    for k, cell_texts in enumerate(texts):
        words = [w for t in cell_texts for w in t.split()]
        need = (max((len(w) for w in words), default=0) * char) + CELL_PADDING_PT
        longest = max((len(t) for t in cell_texts), default=0)
        prose = any(len(t.split()) >= 4 for t in cell_texts)
        want = min(longest * char + CELL_PADDING_PT, PROSE_SHARE * TEXT_WIDTH_PT) if prose else 0
        floor.append(min(need, PROSE_SHARE * TEXT_WIDTH_PT))
        wish.append(max(widths[k], floor[-1], want))
    if all(w >= f - 0.5 for w, f in zip(widths, floor)):
        return widths  # every word fits: the table as printed
    room = max(TEXT_WIDTH_PT, sum(widths))
    if sum(wish) <= room:
        return wish
    if sum(floor) >= room:
        return [f * room / sum(floor) for f in floor]
    # every column keeps its floor; the rest of the page is shared out in
    # proportion to what each wanted beyond it
    spare, slack = room - sum(floor), [w - f for w, f in zip(wish, floor)]
    return [f + spare * s / (sum(slack) or 1.0) for f, s in zip(floor, slack)]


def table(doc, header, rows, layout=None, unit_row=None, shift=lambda s: s):
    """A table drawn like last year's: its printed header rows with merged
    cells, its column widths, Arial 9, figures right-aligned and Total rows
    in bold. `rows` are the new order's cells (order_draft._table_rows), in
    the extracted table's columns; without a recorded layout the merged
    one-row header is used."""
    cols = _kept_columns(header, layout, rows)
    if not cols:
        return None
    pos = {j: k for k, j in enumerate(cols)}
    widths_pt = (layout or {}).get("widths") or []
    widths = [widths_pt[j] if j < len(widths_pt) and widths_pt[j] > 0 else 0 for j in cols]
    # a heading spanning several extracted columns of which only one is
    # printed (Table 5.89's "KSERC Approval" over 6 + 81 + 6 pt, its figures
    # in the 6 pt one) gives that column the whole span's width
    for hr in (layout or {}).get("header_rows", []):
        for c in hr:
            span = range(c["col"], c["col"] + c.get("colspan", 1))
            kept = [j for j in span if j in pos]
            if len(kept) == 1:
                k = pos[kept[0]]
                widths[k] = max(widths[k], sum(widths_pt[j] for j in span if j < len(widths_pt)))
    if not any(widths):
        widths = [1.0] * len(cols)
    total = sum(widths) or 1.0
    scale = min(1.0, TEXT_WIDTH_PT / total)
    widths = [max(w * scale, 18.0) for w in widths]

    header_rows = (layout or {}).get("header_rows") or []
    if not header_rows:
        header_rows = [[{"col": j, "colspan": 1, "rowspan": 1, "text": header[j] if j < len(header) else ""}
                        for j in range(len(header))]]
    n_head = len(header_rows) + (1 if unit_row and any((u or "").strip() for u in unit_row) else 0)
    size = TABLE_PT if len(cols) <= 9 else 8
    widths = _balanced_widths(widths, cols, header_rows, rows, size)

    grid = doc.add_table(rows=n_head + len(rows), cols=len(cols))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    # every cell taken once, straight from the table's XML: python-docx's
    # grid.cell(r, c) / column.cells re-read the whole table on each call,
    # which made a 250-table order take over a minute
    matrix = [[_Cell(tc, grid) for tc in tr.tc_lst] for tr in grid._tbl.tr_lst]
    # fixed layout: Word takes the column widths from the grid and the first
    # row, so they needn't be repeated on every cell
    layout_el = OxmlElement("w:tblLayout")
    layout_el.set(qn("w:type"), "fixed")
    grid._tbl.tblPr.append(layout_el)
    for k, gc in enumerate(grid._tbl.tblGrid.gridCol_lst):
        gc.w = Pt(widths[k])
    for k, cell in enumerate(matrix[0]):
        cell.width = Pt(widths[k])
    styles = _table_styles(doc)
    small = None if size == TABLE_PT else size

    # header rows with their merges, in printed columns
    for r, hr in enumerate(header_rows):
        for c in hr:
            span = [pos[j] for j in range(c["col"], c["col"] + c.get("colspan", 1)) if j in pos]
            if not span:
                continue
            last_row = min(r + c.get("rowspan", 1) - 1, len(header_rows) - 1)
            cell = matrix[r][span[0]]
            if last_row != r or span[-1] != span[0]:
                try:
                    cell = cell.merge(matrix[last_row][span[-1]])
                except Exception:
                    pass  # overlapping spans measured off the page: leave unmerged
            if not cell.text:
                _cell_text(cell, shift(c.get("text") or ""), styles["heading"], small)
    if n_head > len(header_rows):
        for j in cols:
            if j < len(unit_row) and (unit_row[j] or "").strip():
                _cell_text(matrix[len(header_rows)][pos[j]], unit_row[j], styles["heading"], small)
    for r in range(n_head):
        _repeat_header(grid.rows[r])

    for i, row in enumerate(rows):
        label = next((c for c in row if (c or "").strip() and not NUMBER_CELL.match(c.strip())), "")
        bold = bool(TOTAL_ROW.match(label.strip()))
        cells = matrix[n_head + i]
        for j in cols:
            text = row[j] if j < len(row) else ""
            plain = re.sub(r'\[\[\s*\]\]|\{\{|\}\}', '', text or "").strip()
            figure = NUMBER_CELL.match(plain) or (text or "").strip() == "[[ ]]"
            _cell_text(cells[pos[j]], text or "", styles["figure" if figure else "text"], small, bold=bold)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return grid
