import io
import docx

import order_draft


def field(**kw):
    base = {"field_label": "Depreciation", "unit": "Rs. Cr", "arr_approved": "175.00", "petition_claimed": "182.74"}
    base.update(kw)
    return base


def test_pending_field_has_no_approved_figure():
    assert order_draft.effective_approved(field(decision="pending")) is None


def test_accepting_takes_the_claim_and_follows_later_edits_to_it():
    f = field()
    order_draft.apply_decision(f, {"decision": "accepted"})
    assert order_draft.effective_approved(f) == "182.74"
    f["petition_claimed"] = "180.00"
    assert order_draft.effective_approved(f) == "180.00"


def test_typing_a_figure_equal_to_the_claim_means_accepted_otherwise_modified():
    f = field()
    order_draft.apply_decision(f, {"commission_approved": "182.74"})
    assert f["decision"] == "accepted"
    order_draft.apply_decision(f, {"commission_approved": "170.00"})
    assert (f["decision"], order_draft.effective_approved(f)) == ("modified", "170.00")
    order_draft.apply_decision(f, {"commission_approved": ""})
    assert f["decision"] == "pending"


def test_disallowed_approves_nothing_and_unknown_decision_is_refused():
    f = field()
    order_draft.apply_decision(f, {"decision": "disallowed"})
    assert order_draft.effective_approved(f) == "0.00"
    try:
        order_draft.apply_decision(f, {"decision": "maybe"})
    except ValueError:
        return
    raise AssertionError("unknown decision accepted")


def test_draft_states_decisions_and_leaves_pending_items_as_placeholders(monkeypatch):
    monkeypatch.setattr(order_draft, "_template_facts", lambda data: ("2", "25.06.2022"))
    decided = field(decision="modified", commission_approved="170.00", reason="Normative basis applied.")
    pending = field(field_label="RoE", arr_approved="116.38", petition_claimed="116.38", decision="pending")
    data = {"sbu": "G", "target_year": "2024-25", "fields": [decided, pending]}
    document = docx.Document(io.BytesIO(order_draft.build_docx(data)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "at Rs. 170.00 Cr as against Rs. 182.74 Cr claimed by KSEB Ltd" in text
    assert "Normative basis applied." in text
    assert "[Commission's decision on RoE is pending" in text


def test_every_field_is_pending_until_the_reviewer_decides():
    # a claim is never accepted by default: the draft would otherwise say the
    # Commission approves items nobody has looked at
    assert order_draft.decision_of(field()) == "pending"
    assert order_draft.effective_approved(field()) is None
    assert order_draft.decision_of(field(petition_claimed=None)) == "pending"
    assert order_draft.decision_of(field(decision="accepted")) == "accepted"


def test_accept_remaining_claims_only_touches_pending_items_with_a_claim():
    fields = [field(), field(decision="disallowed"), field(petition_claimed=None), field(excluded=True)]
    assert order_draft.accept_remaining_claims(fields) == 1
    assert [order_draft.decision_of(f) for f in fields] == ["accepted", "disallowed", "pending", "pending"]
    assert order_draft.effective_approved(fields[0]) == "182.74"


def test_financial_years_and_year_end_dates_move_on_but_other_years_stay():
    assert order_draft.shift_year("for the year 2023-24 as on 31.03.2024") == "for the year 2024-25 as on 31.03.2025"
    assert order_draft.shift_year("Tariff Regulations, 2021 and the Act, 2003") == "Tariff Regulations, 2021 and the Act, 2003"
    assert order_draft.shift_year("FY 2022 – 23 and 2099-00") == "FY 2023 – 24 and 2100-01"


def test_carried_figures_are_highlighted_but_the_clause_number_is_not():
    marked = order_draft._mark_figures("2.5 KSEB Ltd submitted 5721.17 MU against 7,235.74 MU on 25.06.2022 under Section 131.")
    assert marked.startswith("2.5 KSEB Ltd")
    assert "{{5721.17}}" in marked and "{{7,235.74}}" in marked and "{{25.06.2022}}" in marked
    assert "{{131}}" not in marked


def test_sentences_broken_at_a_hyphen_or_after_rs_are_joined_back():
    blocks = [{"type": "heading", "text": "As above, the Commission hereby approves the O&M expenses of SBU-"},
              {"type": "para", "text": "G of KSEB Ltd for the year 2023-24 at Rs. 186.16 Cr as against Rs."},
              {"type": "para", "text": "189.46 Cr claimed by KSEB Ltd."},
              {"type": "para", "text": "2.35 Next clause."}]
    joined = order_draft._join_hyphen_breaks(blocks)
    assert len(joined) == 2
    assert joined[0]["text"].endswith("Rs. 189.46 Cr claimed by KSEB Ltd.") and "SBU-G of KSEB" in joined[0]["text"]


def test_table_columns_are_recognised_from_the_previous_orders_headings():
    roles = order_draft._column_roles(["No", "Particulars", "MYT Order dated 25.06.2022", "Actuals", "TU Sought",
                                       "Difference over approval", "KSERC Approval"])
    assert roles == {2: "myt", 3: "actual", 4: "claim", 5: "difference", 6: "approved"}


def test_table_rows_fill_known_heads_and_leave_other_figures_as_placeholders():
    table = {"header": ["No", "Particulars", "MYT Order dated 25.06.2022", "Actuals", "TU Sought"],
             "data_rows": [["1", "Depreciation", "157.02", "159.67", "159.67"], ["2", "Unlisted head", "1.50", "1.50", "1.50"]]}
    fields = order_draft._fields_by_key({"fields": [{"field_label": "Depreciation", "arr_approved": "175.00",
                                                     "petition_claimed": "182.74", "petition_actuals": "772.04"}]})
    rows = order_draft._table_rows(table, fields)
    assert rows[0] == ["1", "Depreciation", "175.00", "772.04", "182.74"]
    assert rows[1] == ["2", "Unlisted head", "[[ ]]", "[[ ]]", "[[ ]]"]


def test_a_figure_quoted_from_a_table_moves_on_to_the_figure_now_in_that_cell():
    old = [["1", "Depreciation", "157.02", "182.74"], ["2", "Return on Equity", "116.38", "116.38"]]
    new = [["1", "Depreciation", "175.00", "190.12"], ["2", "Return on Equity", "120.00", "[[ ]]"]]
    updates = order_draft.figure_updates([(old, new)])
    assert updates == {157.02: 175.0, 182.74: 190.12}  # 116.38 is in two cells, one still blank
    marked = order_draft._mark_figures("2.34 Approved Rs.157.02 Cr against Rs. 182.74 Cr; RoE Rs.116.38 Cr.", updates)
    assert marked == "2.34 Approved Rs.<<175.00>> Cr against Rs. <<190.12>> Cr; RoE Rs.{{116.38}} Cr."


def test_figures_that_now_differ_change_sign_or_sit_in_a_removed_row_are_not_moved_on():
    old = [["a", "1.50", "(-)3.30"], ["b", "1.50", "9.10"], ["c", "4.00", "0.00"]]
    new = [["a", "2.00", "3.30"], ["b", "2.50", "9.20"], None]
    assert order_draft.figure_updates([(old, new)]) == {9.1: 9.2}


def test_a_moved_on_figure_is_printed_the_way_the_old_one_was():
    assert order_draft._format_like(123456.7, "1,05,321.20") == "1,23,456.70"
    assert order_draft._format_like(1234.5, "1,053.79") == "1,234.50"
    assert order_draft._format_like(1998.444, "1,998.4") == "1,998.4"
    assert order_draft._format_like(43.6, "43.66") == "43.60"


def test_a_quoted_regulation_keeps_its_years_but_the_commissions_own_italic_moves_on():
    quote = {"italic": True, "bold": False, "size": 11}
    assert order_draft.is_quotation(quote)
    assert not order_draft.is_quotation({"italic": True, "bold": True, "size": 12})
    assert not order_draft.is_quotation({"italic": True, "size": 12})
    text = "To clarify, the O&M expenses for FY 2022-23 shall be allowed at 4.84%."
    assert order_draft._carried_paragraph(text, {}, "SBU-T", "2024-25", quoted=True) == \
        ["To clarify, the O&M expenses for FY 2022-23 shall be allowed at {{4.84}}%."]
    assert "2023-24" in order_draft._carried_paragraph(text, {}, "SBU-T", "2024-25")[0]
